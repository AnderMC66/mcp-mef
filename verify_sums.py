import openpyxl

wb = openpyxl.load_workbook(r"C:\semestre 2026-B\Contabilidad gerencial y tributacion\cuadros_presupuesto_salud_2026.xlsx", data_only=False)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"*** {name} ***")
    cols = {c: 0 for c in range(2, 8)}
    last = ws.max_row
    if ws.cell(last, 1).value and "TOTAL" in str(ws.cell(last, 1).value).upper():
        for r in range(5, last):
            for c in range(2, 8):
                val = ws.cell(r, c).value
                if isinstance(val, (int, float)):
                    cols[c] += val
        print("  Suma calculada:", [cols[c] for c in range(2, 8)])
        avance = cols[6] / cols[3] * 100 if cols[3] else 0
        print(f"  Avance total: {avance:.2f}%")
