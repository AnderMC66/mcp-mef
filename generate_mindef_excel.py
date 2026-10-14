import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DEST_PATH = r"C:\semestre 2026-B\Contabilidad gerencial y tributacion\Cuadros_Presupuestales_MINDEF_08_Octubre.xlsx"

wb = Workbook()
wb.remove(wb.active)

thin_border = Border(
    left=Side(style='thin', color='D0D7DE'),
    right=Side(style='thin', color='D0D7DE'),
    top=Side(style='thin', color='D0D7DE'),
    bottom=Side(style='thin', color='D0D7DE')
)

total_border = Border(
    left=Side(style='thin', color='B0C4DE'),
    right=Side(style='thin', color='B0C4DE'),
    top=Side(style='medium', color='1F497D'),
    bottom=Side(style='double', color='1F497D')
)

def agregar_hoja(wb, title, headers, rows, total_label="TOTAL"):
    ws = wb.create_sheet(title=title)
    
    # Encabezado principal
    ws.merge_cells("A1:H1")
    top = ws["A1"]
    top.value = f"MEF - CONSULTA AMIGABLE: {title.upper()}"
    top.font = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
    top.fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    top.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Subtítulo de fecha
    ws.merge_cells("A2:H2")
    sub = ws["A2"]
    sub.value = "Datos oficiales del portal MEF Transparencia Económica / Consulta Amigable (Año 2026 - al 08 de octubre)"
    sub.font = Font(name="Segoe UI", size=9, italic=True, color="444444")
    sub.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 18

    ws.row_dimensions[3].height = 6

    # Cabecera
    ws.append(headers)
    h_row = 4
    ws.row_dimensions[h_row].height = 24

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=h_row, column=col_idx)
        cell.font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="244061", end_color="244061", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    num_fmt = "#,##0"
    start_row = 5
    for r_idx, r in enumerate(rows, start=start_row):
        ws.append([r[0], r[1], r[2], r[3], r[4], r[5], r[6], f"=F{r_idx}/C{r_idx}"])
        ws.row_dimensions[r_idx].height = 20
        
        for c in range(1, 9):
            cell = ws.cell(row=r_idx, column=c)
            cell.font = Font(name="Segoe UI", size=9)
            cell.border = thin_border
            if c == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif 2 <= c <= 7:
                cell.number_format = num_fmt
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c == 8:
                cell.number_format = "0.0%"
                cell.alignment = Alignment(horizontal="right", vertical="center")

    end_row = start_row + len(rows) - 1
    if total_label:
        t_row = end_row + 1
        total_cells = [
            total_label,
            f"=SUM(B{start_row}:B{end_row})",
            f"=SUM(C{start_row}:C{end_row})",
            f"=SUM(D{start_row}:D{end_row})",
            f"=SUM(E{start_row}:E{end_row})",
            f"=SUM(F{start_row}:F{end_row})",
            f"=SUM(G{start_row}:G{end_row})",
            f"=F{t_row}/C{t_row}"
        ]
        ws.append(total_cells)
        ws.row_dimensions[t_row].height = 24
        for c in range(1, 9):
            cell = ws.cell(row=t_row, column=c)
            cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="1F497D")
            cell.fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
            cell.border = total_border
            if c == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif 2 <= c <= 7:
                cell.number_format = num_fmt
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c == 8:
                cell.number_format = "0.0%"
                cell.alignment = Alignment(horizontal="right", vertical="center")

    ws.column_dimensions["A"].width = 46
    for c_idx in range(2, 8):
        ws.column_dimensions[get_column_letter(c_idx)].width = 17
    ws.column_dimensions["H"].width = 13
    
    ws.freeze_panes = "A5"

headers = [
    "Descripción / Clasificador",
    "PIA (S/)",
    "PIM (S/)",
    "Certificación (S/)",
    "Comp. Anual (S/)",
    "Devengado (S/)",
    "Girado (S/)",
    "% Avance"
]

# 1. Defensa en Ppto Nacional
c1_rows = [
    ["TOTAL PRESUPUESTO NACIONAL", 257561619143, 289196038201, 264424781311, 245512842114, 199314670457, 196474389387],
    ["04: FUNCIÓN DEFENSA Y SEGURIDAD NACIONAL", 7180366014, 9608230663, 9186310451, 8613553918, 6963080911, 6928947639],
    ["26: SECTOR DEFENSA (TODOS LOS PLIEGOS)", 10142901394, 13156958942, 12534231922, 11823795164, 9637979683, 9600062012],
    ["026: MINISTERIO DE DEFENSA (MINDEF CENTRAL)", 9634203776, 12833280814, 12360982783, 11691251653, 9534952013, 9497110974]
]
agregar_hoja(wb, "1. Defensa en Ppto Nacional", headers, c1_rows, total_label=None)

# 2. Sector Defensa por Pliego
c2_rows = [
    ["026: M. DE DEFENSA", 9634203776, 12833280814, 12360982783, 11691251653, 9534952013, 9497110974],
    ["006: INSTITUTO NACIONAL DE DEFENSA CIVIL (INDECI)", 477363805, 292650738, 142620125, 102077984, 81249544, 81219111],
    ["025: CENEPRED", 17993867, 17190468, 16981361, 16890494, 11779149, 11776219],
    ["335: AGENCIA DE COMPRAS DE LAS FF.AA. (ACFFAA)", 8849729, 8690274, 8664277, 8621516, 6163049, 6162149],
    ["332: INSTITUTO GEOGRAFICO NACIONAL (IGN)", 4490217, 5146648, 4983375, 4953519, 3835928, 3793559]
]
agregar_hoja(wb, "2. Sector Defensa por Pliego", headers, c2_rows, total_label="TOTAL SECTOR DEFENSA")

# 3. MINDEF por Unidad Ejecutora (Fuerzas Armadas)
c3_rows = [
    ["003-470: EJERCITO PERUANO", 2948281901, 3547561848, 3332432955, 3184607168, 2355230869, 2344506794],
    ["004-471: MARINA DE GUERRA DEL PERU", 2485582296, 2840463688, 2728423066, 2599974611, 2017849048, 2015819707],
    ["005-472: FUERZA AEREA DEL PERU", 2008494812, 3858849941, 3732204619, 3358426975, 2931546258, 2906928871],
    ["009-1124: OFICINA PREVISIONAL DE LAS FF.AA. (CPMP)", 1757819492, 2386521998, 2386519329, 2377723596, 2102246306, 2102224512],
    ["001-468: MINISTERIO DE DEFENSA-OGA", 115226632, 88676108, 77861213, 74039829, 56196463, 56071830],
    ["002-469: COMANDO CONJUNTO DE LAS FUERZAS ARMADAS", 277242264, 69908759, 63470516, 56794905, 40916623, 40909583],
    ["006-1122: CONIDA (AGENCIA ESPACIAL)", 23184199, 24171361, 23101694, 22834821, 18053342, 17803877],
    ["008-1123: ESCUELA NACIONAL DE MARINA MERCANTE", 10066933, 9877435, 9786120, 9670971, 7587155, 7579497],
    ["011-1761: CAEN-EPG (CENTRO ALTOS ESTUDIOS)", 8305247, 7249676, 7183272, 7178776, 5325950, 5266303]
]
agregar_hoja(wb, "3. MINDEF por Ejecutora", headers, c3_rows, total_label="TOTAL MINDEF")

# 4. MINDEF por Fuente de Financiamiento
c4_rows = [
    ["1: RECURSOS ORDINARIOS", 9203447658, 10431236348, 10014644958, 9563242408, 7474079754, 7436578730],
    ["3: REC. OPERACIONES OFICIALES DE CREDITO", 242917000, 1902270665, 1873171524, 1701128528, 1691892873, 1691892873],
    ["5: RECURSOS DETERMINADOS (CANON / REGALIAS)", 130688515, 368515457, 365830208, 357224721, 311075243, 310905045],
    ["4: DONACIONES Y TRANSFERENCIAS", 30798320, 104906061, 90660970, 56731793, 47596158, 47464455],
    ["2: RECURSOS DIRECTAMENTE RECAUDADOS", 26352283, 26352283, 16675122, 12924203, 10307985, 10269870]
]
agregar_hoja(wb, "4. MINDEF por Fuente", headers, c4_rows, total_label="TOTAL MINDEF")

# 5. MINDEF por Genérica de Gasto
c5_rows = [
    ["5-21: PERSONAL Y OBLIGACIONES SOCIALES", 5248403697, 5293066966, 5039475602, 5017059778, 3743658889, 3732299322],
    ["6-26: ADQUISICION DE ACTIVOS NO FINANCIEROS", 694637090, 2681888671, 2648151663, 2458706595, 2398372679, 2398087512],
    ["5-23: BIENES Y SERVICIOS", 2015392809, 2512603773, 2335258017, 1879039155, 1355748871, 1329572528],
    ["5-22: PENSIONES Y OTRAS PRESTACIONES SOCIALES", 1589203756, 1560307588, 1559650595, 1559454124, 1301571216, 1301551254],
    ["5-25: OTROS GASTOS", 25204194, 723189394, 716349282, 714894377, 701986422, 701986422],
    ["7-28: SERVICIO DE LA DEUDA PUBLICA", 60876800, 60876804, 60876801, 60876801, 32393113, 32393113],
    ["6-24: DONACIONES Y TRANSFERENCIAS CAPITAL", 0, 862188, 862188, 862188, 862188, 862188],
    ["5-24: DONACIONES Y TRANSFERENCIAS", 485430, 485430, 358636, 358636, 358636, 358636]
]
agregar_hoja(wb, "5. MINDEF por Generica", headers, c5_rows, total_label="TOTAL MINDEF")

# 6. MINDEF por Programa Presupuestal
c6_rows = [
    ["0135: CAPACIDADES MILITARES PARA DEFENSA Y DESARROLLO", 7207181471, 9624379722, 9238493951, 8747645248, 7092853835, 7058468959],
    ["9002: ASIGNACIONES PPTALES QUE NO RESULTAN EN PRODUCTOS", 1933075684, 2595610986, 2558503798, 2524097201, 2181949296, 2181804981],
    ["9001: ACCIONES CENTRALES", 207527823, 203328534, 195679801, 189247574, 140703954, 140289325],
    ["0032: LUCHA CONTRA EL TERRORISMO", 174372057, 175540690, 170667487, 116043896, 63218019, 61916016],
    ["0128: REDUCCION DE LA MINERIA ILEGAL", 92266732, 91734063, 88394241, 71343548, 43261162, 41903622],
    ["0068: REDUCCION DE VULNERABILIDAD ANTE DESASTRES", 11664072, 137562337, 104409510, 38134841, 8974834, 8947862],
    ["0137: DESARROLLO DE CIENCIA Y TECNOLOGIA", 2040937, 4059681, 3769401, 3700968, 2979107, 2777617],
    ["0074: CONTROL DE OFERTA DE DROGAS", 0, 989801, 989595, 963576, 963576, 963576],
    ["0031: REDUCCION DEL TRAFICO ILICITO DE DROGAS", 6075000, 75000, 75000, 74800, 48232, 39016]
]
agregar_hoja(wb, "6. MINDEF por Programa", headers, c6_rows, total_label="TOTAL MINDEF")

wb.save(DEST_PATH)
print("Archivo MINDEF generado con éxito en:", DEST_PATH)
