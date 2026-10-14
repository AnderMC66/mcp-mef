import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DEST_PATH = r"C:\semestre 2026-B\Contabilidad gerencial y tributacion\cuadros_presupuesto_salud_2026.xlsx"

wb = Workbook()
wb.remove(wb.active)  # Quitar hoja por defecto

thin_border = Border(
    left=Side(style='thin', color='DDDDDD'),
    right=Side(style='thin', color='DDDDDD'),
    top=Side(style='thin', color='DDDDDD'),
    bottom=Side(style='thin', color='DDDDDD')
)

total_border = Border(
    left=Side(style='thin', color='B0C4DE'),
    right=Side(style='thin', color='B0C4DE'),
    top=Side(style='medium', color='1F497D'),
    bottom=Side(style='double', color='1F497D')
)

def agregar_hoja(wb, title, headers, rows, total_label="TOTAL"):
    ws = wb.create_sheet(title=title)
    
    # 1. Encabezado principal
    ws.merge_cells("A1:H1")
    top = ws["A1"]
    top.value = f"MEF - CONSULTA AMIGABLE: {title.upper()}"
    top.font = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
    top.fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    top.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # 2. Subtítulo con fecha exacta de corte
    ws.merge_cells("A2:H2")
    sub = ws["A2"]
    sub.value = "Datos oficiales del portal MEF Transparencia Económica / Consulta Amigable (Año 2026 - al 08 de octubre)"
    sub.font = Font(name="Segoe UI", size=9, italic=True, color="444444")
    sub.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 18

    # Fila 3 vacía
    ws.row_dimensions[3].height = 6

    # 4. Cabeceras de columnas
    ws.append(headers)
    h_row = 4
    ws.row_dimensions[h_row].height = 24

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=h_row, column=col_idx)
        cell.font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="244061", end_color="244061", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    num_fmt = "#,##0"
    start_row = 5
    for r_idx, r in enumerate(rows, start=start_row):
        ws.append([r[0], r[1], r[2], r[3], r[4], r[5], r[6], f"=F{r_idx}/C{r_idx}"])
        ws.row_dimensions[r_idx].height = 20
        
        for c in range(1, 9):
            cell = ws.cell(row=r_idx, column=c)
            cell.font = Font(name="Segoe UI", size=9)
            cell.border = thin_border
            if c == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif 2 <= c <= 7:
                cell.number_format = num_fmt
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c == 8:
                cell.number_format = "0.0%"
                cell.alignment = Alignment(horizontal="right", vertical="center")

    # Fila de TOTAL si no es Cuadro 1
    end_row = start_row + len(rows) - 1
    if total_label:
        t_row = end_row + 1
        total_cells = [
            total_label,
            f"=SUM(B{start_row}:B{end_row})",
            f"=SUM(C{start_row}:C{end_row})",
            f"=SUM(D{start_row}:D{end_row})",
            f"=SUM(E{start_row}:E{end_row})",
            f"=SUM(F{start_row}:F{end_row})",
            f"=SUM(G{start_row}:G{end_row})",
            f"=F{t_row}/C{t_row}"
        ]
        ws.append(total_cells)
        ws.row_dimensions[t_row].height = 24
        for c in range(1, 9):
            cell = ws.cell(row=t_row, column=c)
            cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="1F497D")
            cell.fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
            cell.border = total_border
            if c == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif 2 <= c <= 7:
                cell.number_format = num_fmt
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c == 8:
                cell.number_format = "0.0%"
                cell.alignment = Alignment(horizontal="right", vertical="center")

    # Ajuste de ancho de columnas
    ws.column_dimensions["A"].width = 44
    for c_idx in range(2, 8):
        ws.column_dimensions[get_column_letter(c_idx)].width = 17
    ws.column_dimensions["H"].width = 13
    
    ws.freeze_panes = "A5"

headers = [
    "Descripción / Clasificador",
    "PIA (S/)",
    "PIM (S/)",
    "Certificación (S/)",
    "Comp. Anual (S/)",
    "Devengado (S/)",
    "Girado (S/)",
    "% Avance"
]

# ==========================================
# 1. Salud en Ppto Nacional (al 08/10/2026)
# ==========================================
c1_rows = [
    ["TOTAL PRESUPUESTO NACIONAL", 257561619143, 289196038201, 264424781311, 245512842114, 199314670457, 196474389387],
    ["20: FUNCIÓN SALUD (TODO EL ESTADO)", 32887683068, 35843460554, 32872305230, 31079920004, 25720861651, 25438605096],
    ["011: MINISTERIO DE SALUD (MINSA CENTRAL)", 11396612777, 11311608497, 10398817132, 10152497304, 8292021159, 8244392583]
]
agregar_hoja(wb, "1. Salud en Ppto Nacional", headers, c1_rows, total_label=None)

# ==========================================
# 2. Función Salud por Nivel (al 08/10/2026)
# ==========================================
c2_rows = [
    ["E: GOBIERNO NACIONAL", 16349450831, 16401206851, 15320757434, 14962958708, 12390374978, 12346326528],
    ["R: GOBIERNOS REGIONALES", 15750399090, 18208730896, 16543608651, 15304595812, 12665971342, 12447419641],
    ["M: GOBIERNOS LOCALES", 787833147, 1233522807, 1007939145, 812365484, 664515331, 644858927]
]
agregar_hoja(wb, "2. Funcion Salud por Nivel", headers, c2_rows, total_label="TOTAL FUNCIÓN SALUD")

# ==========================================
# 3. Sector Salud por Pliego (al 08/10/2026)
# ==========================================
c3_rows = [
    ["011: M. DE SALUD (MINSA)", 11396612777, 11311608497, 10398817132, 10152497304, 8292021159, 8244392583],
    ["135: SEGURO INTEGRAL DE SALUD (SIS)", 2686113352, 2685895168, 2650850363, 2617420688, 2515704541, 2512032864],
    ["136: INST. NAC. ENFERMEDADES NEOPLASICAS (INEN)", 351952428, 534466876, 514256530, 489532416, 380228279, 378779546],
    ["131: INSTITUTO NACIONAL DE SALUD (INS)", 201932107, 202520374, 189760857, 171103489, 123558399, 123252977],
    ["134: SUPERINTENDENCIA NACIONAL DE SALUD (SUSALUD)", 69844662, 69805009, 67512897, 66664463, 48345276, 48179133]
]
agregar_hoja(wb, "3. Sector Salud por Pliego", headers, c3_rows, total_label="TOTAL SECTOR SALUD")

# ==========================================
# 4. MINSA por Fuente (al 08/10/2026 - Cuadra con MINSA Central)
# ==========================================
c4_rows = [
    ["00: RECURSOS ORDINARIOS", 11190571059, 10316207393, 9544736079, 9365837366, 7622031774, 7580248198],
    ["13: DONACIONES Y TRANSFERENCIAS", 0, 733163499, 665524017, 605149794, 519186246, 514417955],
    ["09: RECURSOS DIRECTAMENTE RECAUDADOS", 206041718, 222100903, 171577470, 152446622, 122303751, 121227042],
    ["19: REC. POR OPERACIONES OFICIALES DE CREDITO", 0, 40000000, 29011847, 29011847, 28459201, 28459201],
    ["18: CANON Y SOBRECANON, REGALIAS (RECURSOS DETERMINADOS)", 0, 136702, 51675, 51675, 40187, 40187]
]
agregar_hoja(wb, "4. MINSA por Fuente", headers, c4_rows, total_label="TOTAL MINSA")

# ==========================================
# 5. MINSA por Genérica (al 08/10/2026 - Cuadra con MINSA Central)
# ==========================================
c5_rows = [
    ["5-21: PERSONAL Y OBLIGACIONES SOCIALES", 6334443545, 5304280199, 5087109971, 5017338339, 3847236688, 3837216830],
    ["5-23: BIENES Y SERVICIOS", 2596098674, 3478273194, 2987018750, 2820152973, 2285780141, 2263219871],
    ["6-26: ADQUISICION DE ACTIVOS NO FINANCIEROS", 1591509920, 1639899993, 1582306416, 1566904963, 1424418660, 1410391919],
    ["5-24: DONACIONES Y TRANSFERENCIAS", 584576903, 678304282, 564666561, 558465488, 551983445, 551983445],
    ["5-25: OTROS GASTOS", 160863863, 114947450, 101100808, 100951012, 96951089, 96944284],
    ["5-22: PENSIONES Y OTRAS PRESTACIONES SOCIALES", 129119872, 94903379, 87698582, 87684529, 84651136, 83636234],
    ["6-24: DONACIONES Y TRANSFERENCIAS CAPITAL", 0, 1000000, 1000000, 1000000, 1000000, 1000000]
]
agregar_hoja(wb, "5. MINSA por Generica", headers, c5_rows, total_label="TOTAL MINSA")

# ==========================================
# 6. MINSA por Programa (al 08/10/2026 - Cuadra con MINSA Central)
# ==========================================
c6_rows = [
    ["9002: ASIGNACIONES PPTALES QUE NO RESULTAN EN PRODUCTOS", 5378081454, 5774645796, 5310701591, 5153593329, 4288931672, 4250191732],
    ["9001: ACCIONES CENTRALES", 3758588615, 2484107402, 2351171687, 2313899275, 1750044276, 1745563002],
    ["0104: REDUCCION DE LA MORTALIDAD POR EMERGENCIAS Y URGENCIAS MEDICAS", 280360696, 521821552, 502739087, 498198598, 415838262, 415416396],
    ["0002: SALUD MATERNO NEONATAL", 399906097, 566105663, 507641960, 491555531, 415336643, 414798720],
    ["0024: PREVENCION Y CONTROL DEL CANCER", 495321079, 560143227, 538531808, 516484019, 447127428, 446733374],
    ["1001: PRODUCTOS ESPECIFICOS PARA DESARROLLO INFANTIL TEMPRANO", 440715399, 527680149, 475536179, 470821225, 412223453, 412042171],
    ["0016: TBC-VIH/SIDA", 250841875, 285809548, 207287355, 201537581, 150735001, 148568931],
    ["0018: ENFERMEDADES NO TRANSMISIBLES", 112798380, 201918756, 184769019, 181618809, 147473231, 147390542],
    ["0131: CONTROL Y PREVENCION EN SALUD MENTAL", 113699227, 174359699, 148313039, 142501608, 113595777, 113387112],
    ["0129: PREVENCION Y MANEJO DE CONDICIONES SECUNDARIAS EN DISCAPACIDAD", 63444553, 86631613, 83214796, 82861909, 67370868, 67284773],
    ["0017: ENFERMEDADES METAXENICAS Y ZOONOSIS", 73783053, 87298811, 64227185, 63716301, 55127542, 54848942],
    ["0068: REDUCCION DE VULNERABILIDAD Y ATENCION DE EMERGENCIAS POR DESASTRES", 15020454, 23421497, 19781617, 19438258, 15147331, 15100807],
    ["1002: PRODUCTOS ESPECIFICOS PARA REDUCCION DE LA VIOLENCIA CONTRA LA MUJER", 14051895, 17664784, 16985765, 16270861, 13069675, 13066081]
]
agregar_hoja(wb, "6. MINSA por Programa", headers, c6_rows, total_label="TOTAL MINSA")

wb.save(DEST_PATH)
print("Archivo generado exitosamente en:", DEST_PATH)
