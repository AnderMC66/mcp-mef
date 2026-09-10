"""Lógica de negocio del MCP-MEF, compartida entre el servidor MCP (server.py) y el CLI (cli.py).

Estas funciones no son tools de MCP -- son las implementaciones puras. server.py las expone
como @mcp.tool() y cli.py las expone como subcomandos de terminal.
"""
import asyncio
import csv
import itertools
import json
import os
import re
import sqlite3
import urllib.parse
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

import httpx
import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt  # noqa: E402
from fpdf import FPDF  # noqa: E402
from jinja2 import Template  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.cell import WriteOnlyCell  # noqa: E402
from openpyxl.styles import Font  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

# Directorio de datos del usuario: fijo, independiente de dónde viva el código instalado.
# Así el CLI global ("mef ...") y el servidor MCP para Claude Desktop comparten la misma base
# de datos y los mismos informes, sin importar cómo se invoque cada uno.
APP_DIR = Path(os.environ.get("MCP_MEF_HOME", str(Path.home() / ".mcp-mef")))
DB_PATH = str(APP_DIR / "mef_data.db")
OUTPUT_DIR = str(APP_DIR / "output")

try:
    os.makedirs(OUTPUT_DIR, exist_ok=True)
except OSError as _e:  # pragma: no cover - depende del sistema de archivos
    # Un traceback crudo aquí deja al cliente MCP sin pistas de qué configurar.
    raise SystemExit(
        f"No se pudo crear el directorio de datos '{APP_DIR}': {_e}. "
        "Apunta MCP_MEF_HOME a una ruta con permisos de escritura."
    )

TABLE_NAME_RE = re.compile(r'^[A-Za-z_][A-Za-z0-9_]*$')

# Tope por defecto de filas volcadas a un informe. Un SELECT sin LIMIT sobre un dataset de
# 50k filas produce un PDF inmanejable (y tarda minutos), así que se trunca y se avisa.
MAX_REPORT_ROWS = 5000

API_BASE = "https://api.datosabiertos.mef.gob.pe/DatosAbiertos/v1/datastore_search"
USER_AGENT = "mcp-mef (+https://github.com/rodrigocapiz67-tech/mcp-mef)"
HTTP_TIMEOUT = 120.0

# Medido contra la API: una página de 20 000 filas del dataset de Gasto (63 columnas) pesa
# ~55 MB y tarda ~7 s, mientras que 20 páginas de 1 000 tardan el doble. Por encima de esto
# la API empieza a cortar. `filters`, `q` y `fields` de CKAN NO funcionan en este servidor
# (se cuelgan o se ignoran), así que filtrar sólo es posible después, ya en SQLite.
MAX_PAGE_SIZE = 20000

# Prefijo de la tabla temporal donde se vuelca una descarga antes de reemplazar a la real.
# Empieza por '_meta' para que quede fuera de mef_get_schema y de los nombres que el usuario
# puede crear.
STAGING_PREFIX = "_meta_stg_"

CHART_TYPES = ("bar", "pie", "line")

# Tope de filas que sql_mef_db devuelve por defecto. Medido sobre una tabla de 1 millón de
# filas: un `SELECT *` sin tope tardaba 113 s, ocupaba 1,6 GB de RAM y producía 362 MB de
# JSON — suficiente para tumbar la sesión del cliente MCP. Con fetchmany son 10 ms.
MAX_SQL_ROWS = 1000

# Un gráfico con más puntos que esto es ilegible, y matplotlib tarda minutos en dibujarlo.
MAX_CHART_POINTS = 2000

# Límite duro del formato .xlsx (1 048 576 filas contando la cabecera).
EXCEL_MAX_ROWS = 1_048_575

# Filas por lote al exportar: mantiene la memoria plana sin penalizar la velocidad.
STREAM_BATCH = 10000

# Acciones que el authorizer de SQLite permite en las rutas de sólo lectura (informes,
# gráficos, exportaciones). Cualquier otra (INSERT, UPDATE, DELETE, DROP, ATTACH...) la
# rechaza el propio motor, sin depender de inspeccionar el texto de la consulta.
_READ_ONLY_ACTIONS = frozenset({
    sqlite3.SQLITE_SELECT,
    sqlite3.SQLITE_READ,
    sqlite3.SQLITE_FUNCTION,
    sqlite3.SQLITE_RECURSIVE,
})


# ---------------------------------------------------------------------------
# Helpers internos
# ---------------------------------------------------------------------------

@contextmanager
def _db(read_only: bool = False):
    """Abre la base local y garantiza el cierre incluso si la consulta falla.

    Sin esto cada excepción dejaba una conexión abierta: en Windows eso mantiene el archivo
    .db bloqueado y las escrituras siguientes fallan con 'database is locked'.
    """
    conn = sqlite3.connect(DB_PATH)
    try:
        if read_only:
            conn.set_authorizer(
                lambda action, *_: sqlite3.SQLITE_OK if action in _READ_ONLY_ACTIONS else sqlite3.SQLITE_DENY
            )
        yield conn
    finally:
        conn.close()


def _quote_ident(name: str) -> str:
    """Cita un identificador SQL. Los nombres de columna vienen de la API del MEF, así que
    pueden traer comillas o espacios; duplicarlas evita romper (o inyectar) el DDL."""
    return '"' + str(name).replace('"', '""') + '"'


def _validate_table_name(name: str) -> None:
    if not TABLE_NAME_RE.match(name):
        raise ValueError(
            f"Nombre de tabla inválido: '{name}'. Usa solo letras, números y guion bajo, sin empezar con número."
        )
    lowered = name.lower()
    if lowered.startswith("_meta") or lowered.startswith("sqlite_"):
        raise ValueError("Los nombres de tabla que empiezan con '_meta' o 'sqlite_' están reservados.")


def _slugify(text: str) -> str:
    text = re.sub(r'[^\w\s-]', '', str(text), flags=re.UNICODE).strip().lower()
    text = re.sub(r'[-\s]+', '_', text)
    return text or "reporte"


def _output_path(filename: str, extension: str, default: str = "reporte") -> str:
    """Ruta segura dentro de OUTPUT_DIR.

    El slug elimina barras y puntos, de modo que un nombre como '../../algo.csv' no puede
    escribir fuera de la carpeta de salida.
    """
    stem = str(filename).strip()
    if stem.lower().endswith(extension):
        stem = stem[: -len(extension)]
    return os.path.join(OUTPUT_DIR, f"{_slugify(stem) if stem.strip() else default}{extension}")


def _is_read_query(query: str) -> bool:
    """True si la consulta sólo lee. Acepta CTEs (`WITH ... SELECT`), que la comprobación
    ingenua `startswith('SELECT')` rechazaba."""
    stripped = re.sub(r'\A(?:\s|--[^\n]*\n?|/\*.*?\*/)+', '', query, flags=re.DOTALL)
    return bool(re.match(r'(?i)(select|with|values)\b', stripped))


def _run_read_query(
    query: str, max_rows: int | None = None
) -> tuple[list[str], list[tuple], bool, dict | None]:
    """Ejecuta una consulta de lectura sobre la base local.

    Devuelve (columnas, filas, hubo_truncado, info_de_fuente). Con `max_rows` pide una fila
    de más para saber si había continuación, sin materializar el resto: sobre tablas de
    millones de filas la diferencia con fetchall() es de milisegundos frente a minutos.
    El authorizer de SQLite garantiza que la consulta no pueda escribir aunque el texto
    parezca inofensivo.
    """
    with _db(read_only=True) as conn:
        cursor = conn.cursor()
        cursor.execute(query)
        columns = [d[0] for d in cursor.description] if cursor.description else []
        if max_rows is None:
            rows, truncated = cursor.fetchall(), False
        else:
            rows = cursor.fetchmany(max_rows + 1)
            truncated = len(rows) > max_rows
            rows = rows[:max_rows]
    return columns, rows, truncated, _lookup_source_info(query)


@contextmanager
def _stream_read_query(query: str, batch_size: int = STREAM_BATCH):
    """Recorre una consulta de lectura por lotes, sin cargarla entera en memoria.

    Cede (columnas, generador_de_filas). Exportar un millón de filas con fetchall() pedía
    1,2 GB de RAM; por lotes la memoria se mantiene plana.
    """
    with _db(read_only=True) as conn:
        cursor = conn.cursor()
        cursor.execute(query)
        columns = [d[0] for d in cursor.description] if cursor.description else []

        def _rows():
            while True:
                lote = cursor.fetchmany(batch_size)
                if not lote:
                    return
                yield from lote

        yield columns, _rows()


def _infer_sql_type(values: list[str]) -> str:
    non_empty = [v for v in values if v not in (None, "")]
    if not non_empty:
        return "TEXT"

    # Los códigos con ceros a la izquierda (ubigeo, códigos de entidad, etc.) deben
    # conservarse como TEXT o se pierde información al convertir a número.
    if any(len(v) > 1 and v[0] == '0' and v.lstrip('-').isdigit() for v in non_empty):
        return "TEXT"

    def _is_int(v: str) -> bool:
        try:
            int(v)
            return True
        except ValueError:
            return False

    def _is_float(v: str) -> bool:
        try:
            f = float(v)
        except ValueError:
            return False
        # 'nan' e 'inf' son float() válidos pero no son datos numéricos: si los tratamos como
        # REAL, una columna de texto con "NaN" se convierte silenciosamente en nulos raros.
        return f == f and f not in (float("inf"), float("-inf"))

    if all(_is_int(v) for v in non_empty):
        return "INTEGER"
    if all(_is_float(v) for v in non_empty):
        return "REAL"
    return "TEXT"


def _convert_value(v, sql_type: str):
    if v is None or v == "":
        return None
    try:
        if sql_type == "INTEGER":
            return int(v)
        if sql_type == "REAL":
            return float(v)
    except (ValueError, TypeError):
        return str(v)
    return str(v)


def _ordered_columns(records: list[dict]) -> list[str]:
    """Unión ordenada de las claves de todos los registros.

    Antes se usaban sólo las claves del primer registro: si la API omitía un campo vacío en
    ese registro concreto, la columna desaparecía para todo el dataset.
    """
    columns: dict[str, None] = {}
    for record in records:
        for key in record:
            columns.setdefault(str(key), None)
    return list(columns)


def _register_download(cursor: sqlite3.Cursor, table_name: str, resource_id: str, row_count: int) -> None:
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS _meta_downloads (
            table_name TEXT PRIMARY KEY,
            resource_id TEXT,
            source_url TEXT,
            fetched_at TEXT,
            row_count INTEGER
        )
    ''')
    source_url = f"{API_BASE}?resource_id={resource_id}"
    cursor.execute('''
        INSERT INTO _meta_downloads (table_name, resource_id, source_url, fetched_at, row_count)
        VALUES (?, ?, ?, ?, ?)
        ON CONFLICT(table_name) DO UPDATE SET
            resource_id=excluded.resource_id,
            source_url=excluded.source_url,
            fetched_at=excluded.fetched_at,
            row_count=excluded.row_count
    ''', (table_name, resource_id, source_url, datetime.now().isoformat(timespec="seconds"), row_count))


def _resolve_resource_id(value: str) -> str:
    """Traduce un alias del catálogo a su resource_id; si no lo es, lo devuelve tal cual.

    Así `mef fetch gasto_2026 tabla` funciona igual que con el UUID, que es lo que hacía útil
    registrar alias en primer lugar.
    """
    candidato = str(value).strip()
    try:
        with _db() as conn:
            row = conn.execute(
                "SELECT resource_id FROM _meta_catalog WHERE alias = ?", (candidato,)
            ).fetchone()
    except sqlite3.Error:
        return candidato
    return row[0] if row else candidato


def _user_tables(conn: sqlite3.Connection) -> list[str]:
    """Tablas de datos, dejando fuera las internas (_meta*) y las del propio SQLite."""
    filas = conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' "
        "AND name NOT LIKE '\\_meta%' ESCAPE '\\' AND name NOT LIKE 'sqlite_%' ORDER BY name"
    ).fetchall()
    return [f[0] for f in filas]


def _lookup_source_info(query: str) -> dict | None:
    match = re.search(r'FROM\s+"?([A-Za-z_][A-Za-z0-9_]*)"?', query, re.IGNORECASE)
    if not match:
        return None
    table = match.group(1)
    try:
        with _db() as conn:
            row = conn.execute(
                "SELECT resource_id, source_url, fetched_at FROM _meta_downloads WHERE table_name = ?",
                (table,),
            ).fetchone()
        if row:
            return {"tabla": table, "resource_id": row[0], "source_url": row[1], "fetched_at": row[2]}
    except sqlite3.Error:
        pass
    return None


def _human_size(nbytes: float) -> str:
    """Tamaño legible: una base recién creada no debe leerse como '0.0 MB'."""
    for unidad in ("B", "KB", "MB"):
        if abs(nbytes) < 1024:
            return f"{nbytes:.0f} {unidad}" if unidad != "MB" else f"{nbytes:.1f} MB"
        nbytes /= 1024
    return f"{nbytes:.2f} GB"


def _safe_pdf_text(s) -> str:
    # Las fuentes core de fpdf2 (Helvetica) sólo soportan latin-1.
    return str(s).encode('latin-1', 'replace').decode('latin-1')


def _rows_footer(shown: int, truncated: bool) -> str:
    """Pie de un informe.

    Cuando la consulta se cortó en `max_rows` no se conoce el total real (contarlo obligaría
    a recorrer la tabla entera), así que se dice lo que sí se sabe en vez de inventar cifras.
    """
    if not truncated:
        return f"Total de filas: {shown}"
    return f"Filas mostradas: {shown} (la consulta devuelve más; usa LIMIT o sube max_rows)"


def _build_markdown_report(title, columns, rows, generated_at, source_info, truncated) -> str:
    lines = [f"# {title}", "", f"*Generado: {generated_at}*"]
    if source_info:
        lines.append(
            f"*Fuente: resource_id `{source_info['resource_id']}` "
            f"(tabla `{source_info['tabla']}`, descargada {source_info['fetched_at']})*"
        )
    lines.append("")
    lines.append("| " + " | ".join(str(c) for c in columns) + " |")
    lines.append("|" + "|".join(["---"] * len(columns)) + "|")
    for r in rows:
        # Un '|' dentro de una celda parte la tabla Markdown en columnas fantasma.
        cells = ["" if v is None else str(v).replace("|", "\\|").replace("\n", " ") for v in r]
        lines.append("| " + " | ".join(cells) + " |")
    lines.append("")
    lines.append(f"_{_rows_footer(len(rows), truncated)}_")
    return "\n".join(lines)


_HTML_REPORT_TEMPLATE = Template("""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{{ title }}</title>
<style>
  body { font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; padding: 40px; background:#f4f7f6; color:#333; }
  .container { max-width: 1100px; margin:0 auto; background:#fff; padding:30px; border-radius:8px; box-shadow:0 4px 6px rgba(0,0,0,.1); }
  h1 { color:#2c3e50; border-bottom:2px solid #3498db; padding-bottom:10px; }
  .table-wrap { overflow-x:auto; }
  table { width:100%; border-collapse:collapse; margin-top:15px; }
  th, td { padding:10px; text-align:left; border-bottom:1px solid #ddd; font-size:14px; }
  th { background:#3498db; color:#fff; }
  tr:hover { background:#f1f1f1; }
  .meta { color:#7f8c8d; font-size:12px; margin-top:6px; }
  .footer { text-align:center; margin-top:30px; font-size:12px; color:#7f8c8d; }
</style>
</head>
<body>
<div class="container">
  <h1>{{ title }}</h1>
  <div class="meta">
    Generado: {{ generated_at }}
    {% if source_info %} &middot; Fuente: resource_id {{ source_info.resource_id }} (descargado {{ source_info.fetched_at }}){% endif %}
  </div>
  <div class="table-wrap">
    <table>
      <thead><tr>{% for c in columns %}<th>{{ c }}</th>{% endfor %}</tr></thead>
      <tbody>
      {% for row in rows %}
        <tr>{% for v in row %}<td>{{ v if v is not none else "" }}</td>{% endfor %}</tr>
      {% endfor %}
      </tbody>
    </table>
  </div>
  <div class="footer">{{ footer }} &middot; Generado automáticamente por MCP-MEF</div>
</div>
</body>
</html>
""")


def _build_html_report(title, columns, rows, generated_at, source_info, truncated) -> str:
    return _HTML_REPORT_TEMPLATE.render(
        title=title,
        columns=columns,
        rows=rows,
        generated_at=generated_at,
        source_info=source_info,
        footer=_rows_footer(len(rows), truncated),
    )


def _build_pdf_report(title, columns, rows, generated_at, source_info, path, truncated) -> None:
    # Con muchas columnas la tabla no cabe en vertical; el apaisado evita el desbordamiento.
    pdf = FPDF(orientation="L" if len(columns) > 6 else "P")
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, _safe_pdf_text(title), new_x="LMARGIN", new_y="NEXT")

    meta = f"Generado: {generated_at}"
    if source_info:
        meta += f" | Fuente: resource_id {source_info['resource_id']} (descargado {source_info['fetched_at']})"
    pdf.set_font("Helvetica", "", 9)
    pdf.multi_cell(0, 6, _safe_pdf_text(meta), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    pdf.set_font("Helvetica", "", 8)
    with pdf.table() as table:
        header = table.row()
        for c in columns:
            header.cell(_safe_pdf_text(c))
        for r in rows:
            trow = table.row()
            for v in r:
                trow.cell(_safe_pdf_text("" if v is None else v))

    pdf.set_font("Helvetica", "I", 8)
    pdf.ln(4)
    pdf.cell(
        0, 6,
        _safe_pdf_text(_rows_footer(len(rows), truncated)),
        new_x="LMARGIN", new_y="NEXT",
    )
    pdf.output(path)


def _chart_series(query: str) -> tuple[list[str], list[float], str] | str:
    """Ejecuta la consulta de un gráfico y valida su forma.

    Devuelve (etiquetas, valores, nota) o un mensaje de error listo para mostrar. Las filas
    con valor NULL se omiten en vez de tumbar el gráfico entero: un `SUM(...) GROUP BY` casi
    siempre trae alguna categoría sin datos, y antes eso hacía fallar toda la consulta.
    """
    if not _is_read_query(query):
        return "Error: solo se aceptan consultas de lectura (SELECT / WITH)."
    try:
        columns, rows, truncated, _ = _run_read_query(query, MAX_CHART_POINTS)
    except sqlite3.Error as e:
        return f"Error al ejecutar la consulta: {e}"

    if not rows:
        return "La consulta no retornó datos."
    if len(columns) < 2:
        return "La consulta SQL debe retornar 2 columnas: etiqueta y valor numérico."
    if truncated:
        return (
            f"La consulta devuelve más de {MAX_CHART_POINTS} puntos: un gráfico así es "
            "ilegible. Agrupa con GROUP BY o recorta con LIMIT."
        )

    labels, values, omitidas = [], [], 0
    for label, value in ((r[0], r[1]) for r in rows):
        if value is None:
            omitidas += 1
            continue
        try:
            values.append(float(value))
        except (ValueError, TypeError):
            return "La segunda columna de la consulta SQL debe ser numérica para poder graficar."
        labels.append(str(label))

    if not values:
        return "Todos los valores de la segunda columna son NULL: no hay nada que graficar."
    nota = f" ({omitidas} fila(s) con valor NULL omitidas)" if omitidas else ""
    return labels, values, nota


# ---------------------------------------------------------------------------
# Operaciones (usadas por server.py como tools MCP y por cli.py como comandos)
# ---------------------------------------------------------------------------

async def _fetch_page(
    client: httpx.AsyncClient, resource_id: str, limit: int, offset: int
) -> tuple[list[dict], int | None]:
    """Pide una página al datastore y devuelve (registros, total_de_filas_del_dataset)."""
    response = await client.get(
        API_BASE, params={"resource_id": resource_id, "limit": limit, "offset": offset}
    )
    if response.status_code != 200:
        raise httpx.HTTPStatusError(
            f"HTTP {response.status_code}: {response.text[:300]}",
            request=response.request,
            response=response,
        )

    data = response.json()
    records = data.get("records") or []
    result = data.get("result")
    if not records and isinstance(result, dict):
        records = result.get("records") or []

    total = None
    if isinstance(result, dict):
        try:
            total = int(result.get("include_total"))
        except (TypeError, ValueError):
            total = None
    return records, total


async def mef_dataset_info(resource_id: str) -> str:
    """Consulta el tamaño y las columnas de un dataset SIN descargarlo.

    Conviene llamarla antes de cualquier fetch: los datasets de Transparencia Económica
    (Consulta Amigable) pasan de 7 millones de filas y 63 columnas, y bajarlos a ciegas
    puede tardar horas. Acepta también un alias registrado en el catálogo.
    """
    resource_id = _resolve_resource_id(resource_id)
    try:
        async with httpx.AsyncClient(
            timeout=HTTP_TIMEOUT,
            headers={"User-Agent": USER_AGENT},
            transport=httpx.AsyncHTTPTransport(retries=2),
        ) as client:
            records, total = await _fetch_page(client, resource_id, limit=1, offset=0)
    except Exception as e:
        return f"Excepción al consultar el dataset: {e}"

    if not records:
        return "El recurso no devolvió registros: revisa el resource_id."

    columns = _ordered_columns(records)
    info = {
        "resource_id": resource_id,
        "filas_totales": total,
        "columnas": len(columns),
        "nombres_de_columna": columns,
        "muestra": {c: records[0].get(c) for c in columns},
    }
    if total:
        paginas = -(-total // MAX_PAGE_SIZE)
        info["nota"] = (
            f"Descargarlo entero son ~{paginas} paginas de {MAX_PAGE_SIZE} filas; "
            "usa 'limit' para traer solo una parte."
        )
    return json.dumps(info, indent=2, ensure_ascii=False, default=str)


async def fetch_mef_dataset(
    resource_id: str,
    table_name: str,
    limit: int = 1000,
    page_size: int = 1000,
    offset: int = 0,
    progress=None,
    append: bool = False,
) -> str:
    """Descarga un dataset del MEF a SQLite, página a página.

    Cada página se inserta y se confirma en cuanto llega, sobre una tabla de staging: así la
    memoria no crece con el tamaño del dataset (los de Gasto pasan de 7 millones de filas) y
    la tabla anterior sigue intacta y consultable hasta que la descarga termina bien.
    `resource_id` puede ser el UUID del recurso o un alias registrado en el catálogo.
    `progress`, si se pasa, se llama con (filas_descargadas, filas_totales_o_None).

    Con `append=True` este lote se AGREGA a `table_name` en vez de reemplazarla, lo que
    permite traer un dataset de millones de filas en varias llamadas sin que cada una borre
    la anterior. Si además `offset` se deja en 0, se autodetecta como el nº de filas que la
    tabla ya tiene, así que repetir el mismo comando exacto continúa donde se quedó:

        mef fetch <id> gasto_2026 --append --limit 500000   # se repite hasta agotar el dataset

    Cuando ya no quedan filas por traer, devuelve un aviso de que la descarga terminó (no un
    error). Las columnas nuevas que aparezcan en un lote posterior se agregan a la tabla con
    ALTER TABLE; las columnas que un lote no trae quedan NULL en esas filas.
    """
    try:
        _validate_table_name(table_name)
    except ValueError as e:
        return f"Error: {e}"

    if limit <= 0 or page_size <= 0:
        return "Error: 'limit' y 'page_size' deben ser mayores que cero."
    if offset < 0:
        return "Error: 'offset' no puede ser negativo."

    # Admite tanto el resource_id como un alias ya registrado con mef_register_dataset_alias.
    resource_id = _resolve_resource_id(resource_id)
    page_size = min(page_size, MAX_PAGE_SIZE)

    quoted = _quote_ident(table_name)
    staging = _quote_ident(STAGING_PREFIX + table_name)

    columns: list[str] = []
    col_types: dict[str, str] = {}
    downloaded = 0
    total_api: int | None = None
    warning = ""

    try:
        with _db() as conn:
            conn.isolation_level = None
            cur = conn.cursor()
            cur.execute(f"DROP TABLE IF EXISTS {staging}")  # restos de una descarga anterior

            existe_destino = bool(
                cur.execute(
                    "SELECT 1 FROM sqlite_master WHERE type='table' AND name = ?", (table_name,)
                ).fetchone()
            )
            # offset=0 con append=True significa "continúa donde esta tabla se quedó": evita
            # que el usuario tenga que llevar la cuenta de las filas a mano llamada tras llamada.
            aviso_id = ""
            if append and existe_destino:
                if offset == 0:
                    offset = cur.execute(f"SELECT COUNT(*) FROM {quoted}").fetchone()[0]
                # Evita mezclar por error dos descargas distintas (p.ej. Gasto 2025 sobre
                # Gasto 2026) en una misma tabla: se avisa, pero no se bloquea, porque la
                # tabla también pudo crearse a mano con `mef sql`.
                try:
                    previo = cur.execute(
                        "SELECT resource_id FROM _meta_downloads WHERE table_name = ?", (table_name,)
                    ).fetchone()
                except sqlite3.OperationalError:
                    previo = None
                if previo and previo[0] and previo[0] != resource_id:
                    aviso_id = (
                        f" (aviso: '{table_name}' se descargó antes desde resource_id "
                        f"'{previo[0]}', distinto del actual '{resource_id}')"
                    )

            try:
                async with httpx.AsyncClient(
                    timeout=HTTP_TIMEOUT,
                    headers={"User-Agent": USER_AGENT},
                    transport=httpx.AsyncHTTPTransport(retries=2),
                ) as client:
                    while downloaded < limit:
                        batch_limit = min(page_size, limit - downloaded)
                        try:
                            records, total = await _fetch_page(
                                client, resource_id, batch_limit, offset + downloaded
                            )
                        except Exception as e:
                            if not downloaded:
                                raise
                            # Ya hay datos utilizables: mejor conservarlos y avisar.
                            warning = f" (advertencia: la paginación se detuvo por un error: {e})"
                            break

                        if total is not None:
                            total_api = total
                        if not records:
                            break

                        if not columns:
                            columns = _ordered_columns(records)
                            col_types = {
                                col: _infer_sql_type(
                                    [str(r[col]) if r.get(col) is not None else "" for r in records if col in r]
                                )
                                for col in columns
                            }
                            cols_def = ", ".join(f"{_quote_ident(c)} {col_types[c]}" for c in columns)
                            cur.execute(f"CREATE TABLE {staging} ({cols_def})")
                        else:
                            # Una página posterior puede traer una columna que no venía antes.
                            for col in _ordered_columns(records):
                                if col not in col_types:
                                    columns.append(col)
                                    col_types[col] = "TEXT"
                                    cur.execute(
                                        f"ALTER TABLE {staging} ADD COLUMN {_quote_ident(col)} TEXT"
                                    )

                        col_list = ", ".join(_quote_ident(c) for c in columns)
                        placeholders = ", ".join("?" for _ in columns)
                        rows = [
                            tuple(_convert_value(r.get(c), col_types[c]) for c in columns) for r in records
                        ]
                        cur.execute("BEGIN")
                        cur.executemany(
                            f"INSERT INTO {staging} ({col_list}) VALUES ({placeholders})", rows
                        )
                        cur.execute("COMMIT")

                        downloaded += len(records)
                        if progress is not None:
                            progress(downloaded, total_api)
                        if len(records) < batch_limit:
                            break  # última página

                if not downloaded:
                    cur.execute(f"DROP TABLE IF EXISTS {staging}")
                    if append and existe_destino:
                        total_ya = cur.execute(f"SELECT COUNT(*) FROM {quoted}").fetchone()[0]
                        return (
                            f"Descarga completa: '{table_name}' ya tiene todas las filas "
                            f"disponibles ({total_ya}). No quedaban registros nuevos en offset={offset}."
                        )
                    return "No se encontraron registros o la API del MEF indicó error."

                cur.execute("BEGIN")
                if append and existe_destino:
                    # Fusiona el esquema: columnas que la tabla real no tenía se agregan;
                    # columnas que este lote no trae quedan NULL en las filas nuevas (el
                    # INSERT con lista explícita de columnas ya lo hace por sí solo).
                    existentes = {r[1] for r in cur.execute(f"PRAGMA table_info({quoted})")}
                    for col in columns:
                        if col not in existentes:
                            cur.execute(f"ALTER TABLE {quoted} ADD COLUMN {_quote_ident(col)} {col_types[col]}")
                    col_list = ", ".join(_quote_ident(c) for c in columns)
                    cur.execute(f"INSERT INTO {quoted} ({col_list}) SELECT {col_list} FROM {staging}")
                    cur.execute(f"DROP TABLE {staging}")
                    total_registrado = cur.execute(f"SELECT COUNT(*) FROM {quoted}").fetchone()[0]
                else:
                    # Sustitución atómica: hasta este punto la tabla real (si existía) sigue intacta.
                    cur.execute(f"DROP TABLE IF EXISTS {quoted}")
                    cur.execute(f"ALTER TABLE {staging} RENAME TO {quoted}")
                    total_registrado = downloaded
                _register_download(cur, table_name, resource_id, total_registrado)
                cur.execute("COMMIT")
            except Exception:
                # La tabla real no se tocó (o sólo se le agregaron columnas, inofensivo);
                # sólo hay que limpiar el staging.
                try:
                    cur.execute("ROLLBACK")
                except sqlite3.Error:
                    pass
                try:
                    cur.execute(f"DROP TABLE IF EXISTS {staging}")
                except sqlite3.Error:
                    pass
                raise
    except Exception as e:
        detalle = f" (se habían descargado {downloaded} filas)" if downloaded else ""
        return f"Error al descargar el dataset{detalle}: {e}"

    tipos = ", ".join(f"{c}:{t}" for c, t in col_types.items())
    restante = ""
    if total_api and offset + downloaded < total_api:
        pendientes = total_api - offset - downloaded
        si_append = " (con --append para seguir agregando a la misma tabla)" if not append else ""
        restante = (
            f"\nEl dataset tiene {total_api} filas en total: quedan {pendientes} sin descargar. "
            f"Repite el comando{si_append} para seguir, o sube --limit."
        )
    resumen_tabla = f" ({total_registrado} en total en '{table_name}')" if append and existe_destino else ""
    return (
        f"Éxito: se guardaron {downloaded} registros nuevos{resumen_tabla}{warning}{aviso_id}.\n"
        f"Tipos detectados -> {tipos}{restante}"
    )


# El portal datosabiertos.gob.pe (CKAN clásico) que se usaba antes fue dado de baja.
# El portal propio del MEF (datosabiertos.mef.gob.pe) es hoy una SPA en Angular sin API CKAN
# pública documentada; estas dos rutas son las que usa internamente su propio frontend
# (extraídas de su bundle JS: this.main.ws("/PortalWebDatasets/v1.0/getDatasets", ...) y
# this.main.ws("/PortalWebDatasetDetalle/v1.0/getDatasetDetalle", ...)). No son una API oficial
# documentada, así que podrían cambiar sin aviso si el MEF actualiza su portal.
_MEF_SEARCH_URL = "https://datosabiertos.mef.gob.pe/Rest/PortalWebDatasets/v1.0/getDatasets"
_MEF_DATASET_DETAIL_URL = "https://datosabiertos.mef.gob.pe/Rest/PortalWebDatasetDetalle/v1.0/getDatasetDetalle"

_TABULAR_FORMATS = {"CSV", "JSON", "XLS", "XLSX"}


def _portal_client() -> httpx.AsyncClient:
    return httpx.AsyncClient(
        timeout=HTTP_TIMEOUT,
        headers={"User-Agent": USER_AGENT},
        transport=httpx.AsyncHTTPTransport(retries=2),
    )


async def _portal_search(client: httpx.AsyncClient, query: str, page: int) -> dict:
    payload = {
        "search": query, "sort": None, "tags": [], "formats": [], "categories": [], "page": page
    }
    response = await client.post(_MEF_SEARCH_URL, json=payload)
    response.raise_for_status()
    return response.json()


async def _dataset_detail(client: httpx.AsyncClient, slug: str) -> dict:
    response = await client.post(_MEF_DATASET_DETAIL_URL, json={"dataset": slug})
    response.raise_for_status()
    return response.json()


def _resource_entry(res: dict) -> dict:
    return {
        "recurso": res.get("resource_title"),
        "formato": res.get("resource_format"),
        "resource_id": res.get("resource_id"),
        "url_csv": res.get("resource_url"),
    }


async def mef_search_datasets(query: str, page: int = 1) -> str:
    """Busca datasets en el portal y devuelve sus slugs y una muestra de recursos."""
    try:
        async with _portal_client() as client:
            data = await _portal_search(client, query, max(1, page))
            datasets = data.get("datasets") or []
            usada = query

            # El buscador del portal es un LIKE sobre el título/descripción completos, así que
            # una frase de varias palabras casi nunca casa. Si no hay nada, se reintenta con
            # la palabra más larga, que es la que más discrimina.
            if not datasets and len(query.split()) > 1:
                for palabra in sorted(query.split(), key=len, reverse=True):
                    if len(palabra) < 4:
                        continue
                    data = await _portal_search(client, palabra, 1)
                    if data.get("datasets"):
                        datasets = data["datasets"]
                        usada = palabra
                        break

            if not datasets:
                return f"No se encontraron datasets para '{query}'."

            slugs = [ds.get("dataset_id") for ds in datasets[:6] if ds.get("dataset_id")]
            # Los detalles son llamadas independientes: en paralelo cuestan un viaje, no seis.
            detalles = await asyncio.gather(
                *(_dataset_detail(client, slug) for slug in slugs), return_exceptions=True
            )
    except Exception as e:
        return f"Excepción al buscar dataset: {e}"

    por_slug = {
        slug: det for slug, det in zip(slugs, detalles) if not isinstance(det, BaseException)
    }

    salida = []
    for ds in datasets:
        slug = ds.get("dataset_id")
        entrada = {"slug": slug, "titulo": ds.get("dataset"), "formatos": ds.get("formats")}
        recursos = [
            _resource_entry(r)
            for r in (por_slug.get(slug, {}).get("resources") or [])
            if (r.get("resource_format") or "").upper() in _TABULAR_FORMATS
        ]
        if recursos:
            entrada["recursos"] = recursos[:3]
            if len(recursos) > 3:
                entrada["mas_recursos"] = len(recursos) - 3
        salida.append(entrada)

    resultado = {
        "consulta": usada,
        "total_datasets": data.get("countTotal"),
        "pagina": page,
        "paginas": data.get("pagesCount"),
        "datasets": salida,
        "siguiente_paso": (
            "Usa mef_list_resources(slug) para ver todos los recursos de un dataset "
            "(los de varios años traen uno por año) y mef_dataset_info(resource_id) para "
            "saber cuántas filas tiene antes de descargarlo."
        ),
    }
    if usada != query:
        resultado["nota"] = f"'{query}' no dio resultados; se buscó '{usada}'."
    return json.dumps(resultado, indent=2, ensure_ascii=False)


async def mef_list_resources(dataset_slug: str) -> str:
    """Lista todos los recursos descargables de un dataset del portal.

    Es el paso intermedio para los datasets de Transparencia Económica: publican un recurso
    por año (2009-Gasto.csv, 2026-Gasto-Mensual.csv...), cada uno con su propio resource_id.
    """
    slug = dataset_slug.strip()
    if not slug:
        return "Error: indica el slug del dataset (ej. 'presupuesto-y-ejecucion-de-gasto')."

    try:
        async with _portal_client() as client:
            detalle = await _dataset_detail(client, slug)
    except Exception as e:
        return f"Excepción al consultar el dataset: {e}"

    recursos = detalle.get("resources") or []
    if not recursos:
        return (
            f"El dataset '{slug}' no devolvió recursos. Comprueba el slug con mef_search_datasets "
            "(es el campo 'slug', no el título)."
        )

    info = detalle.get("dataset_detail") or {}
    if isinstance(info, list):
        info = info[0] if info else {}
    return json.dumps(
        {
            "slug": slug,
            "titulo": info.get("dataset") or info.get("title"),
            "total_recursos": len(recursos),
            "recursos": [_resource_entry(r) for r in recursos],
        },
        indent=2,
        ensure_ascii=False,
    )


def mef_register_dataset_alias(alias: str, resource_id: str, description: str = "") -> str:
    alias = alias.strip()
    resource_id = resource_id.strip()
    if not alias or not resource_id:
        return "Error: 'alias' y 'resource_id' no pueden estar vacíos."

    try:
        with _db() as conn:
            cursor = conn.cursor()
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS _meta_catalog (
                    alias TEXT PRIMARY KEY,
                    resource_id TEXT,
                    description TEXT,
                    registered_at TEXT
                )
            ''')
            cursor.execute('''
                INSERT INTO _meta_catalog (alias, resource_id, description, registered_at)
                VALUES (?, ?, ?, ?)
                ON CONFLICT(alias) DO UPDATE SET
                    resource_id=excluded.resource_id,
                    description=excluded.description,
                    registered_at=excluded.registered_at
            ''', (alias, resource_id, description, datetime.now().isoformat(timespec="seconds")))
            conn.commit()
    except sqlite3.Error as e:
        return f"Error al registrar el alias: {e}"

    return f"Alias '{alias}' registrado -> resource_id '{resource_id}'."


def mef_remove_dataset_alias(alias: str) -> str:
    """Borra un alias del catálogo (no toca las tablas ya descargadas)."""
    nombre = alias.strip()
    if not nombre:
        return "Error: indica el alias a borrar."
    try:
        with _db() as conn:
            cursor = conn.execute("DELETE FROM _meta_catalog WHERE alias = ?", (nombre,))
            conn.commit()
            borrados = cursor.rowcount
    except sqlite3.OperationalError:
        return "Aún no hay ningún alias registrado."
    except sqlite3.Error as e:
        return f"Error al borrar el alias: {e}"

    if not borrados:
        return f"No había ningún alias llamado '{nombre}'."
    return f"Alias '{nombre}' borrado."


def mef_list_known_datasets() -> str:
    try:
        with _db() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='_meta_catalog'")
            if not cursor.fetchone():
                return (
                    "Aún no hay datasets registrados. Usa mef_register_dataset_alias tras "
                    "encontrar uno con mef_search_datasets."
                )
            cursor.execute(
                "SELECT alias, resource_id, description, registered_at FROM _meta_catalog ORDER BY registered_at DESC"
            )
            rows = cursor.fetchall()
    except sqlite3.Error as e:
        return f"Error al leer el catálogo: {e}"

    catalog = [
        {"alias": a, "resource_id": r, "descripcion": d, "registrado": t}
        for a, r, d, t in rows
    ]
    return json.dumps(catalog, indent=2, ensure_ascii=False)


def sql_mef_db(query: str, max_rows: int = MAX_SQL_ROWS) -> str:
    """Ejecuta SQL sobre la base local. `max_rows=0` quita el tope (úsalo con cuidado)."""
    try:
        with _db() as conn:
            cursor = conn.cursor()
            cursor.execute(query)
            # `cursor.description` es la señal fiable de si la consulta devolvió filas:
            # cubre CTEs (`WITH ... SELECT`) y PRAGMA, que un startswith('SELECT') no ve.
            if cursor.description is not None:
                columns = [desc[0] for desc in cursor.description]
                if max_rows > 0:
                    results = cursor.fetchmany(max_rows + 1)
                    truncado = len(results) > max_rows
                    results = results[:max_rows]
                else:
                    results, truncado = cursor.fetchall(), False
                payload = {"columnas": columns, "filas": results}
                if truncado:
                    payload["truncado"] = True
                    payload["nota"] = (
                        f"Se devuelven las primeras {max_rows} filas. Agrega LIMIT, agrupa con "
                        "GROUP BY o sube max_rows si de verdad necesitas más."
                    )
                return json.dumps(payload, indent=2, ensure_ascii=False, default=str)
            conn.commit()
            return json.dumps({
                "status": "success",
                "message": "Consulta ejecutada correctamente.",
                "filas_afectadas": cursor.rowcount,
            }, ensure_ascii=False)
    except Exception as e:
        return json.dumps({"status": "error", "message": str(e)}, ensure_ascii=False)


def mef_get_schema() -> str:
    try:
        with _db() as conn:
            cursor = conn.cursor()
            tables = _user_tables(conn)

            schema = {}
            for table in tables:
                quoted = _quote_ident(table)
                cursor.execute(f'PRAGMA table_info({quoted})')
                columns = [row[1] for row in cursor.fetchall()]
                cursor.execute(f'SELECT COUNT(*) FROM {quoted}')
                row_count = cursor.fetchone()[0]
                entry = {"columnas": columns, "filas": row_count}
                try:
                    cursor.execute(
                        "SELECT resource_id, fetched_at FROM _meta_downloads WHERE table_name = ?", (table,)
                    )
                    meta = cursor.fetchone()
                    if meta:
                        entry["resource_id"] = meta[0]
                        entry["descargado"] = meta[1]
                except sqlite3.OperationalError:
                    pass
                schema[table] = entry

        if not schema:
            return (
                "La base local está vacía. Descarga un dataset primero con fetch_mef_dataset "
                "(o `mef fetch <resource_id> <tabla>`)."
            )
        return json.dumps(schema, indent=2, ensure_ascii=False)
    except Exception as e:
        return f"Error al obtener el esquema: {e}"


def mef_db_stats() -> str:
    """Tamaño de la base en disco y qué ocupa cada tabla descargada."""
    try:
        tamano = os.path.getsize(DB_PATH) if os.path.exists(DB_PATH) else 0
        with _db() as conn:
            tablas = {
                nombre: conn.execute(f'SELECT COUNT(*) FROM {_quote_ident(nombre)}').fetchone()[0]
                for nombre in _user_tables(conn)
            }
    except sqlite3.Error as e:
        return f"Error al leer la base: {e}"

    return json.dumps(
        {
            "base_de_datos": DB_PATH,
            "tamano": _human_size(tamano),
            "tamano_mb": round(tamano / 1e6, 2),
            "salidas": OUTPUT_DIR,
            "tablas": tablas,
            "filas_totales": sum(tablas.values()),
        },
        indent=2,
        ensure_ascii=False,
    )


def mef_drop_table(table_name: str, vacuum: bool = False) -> str:
    """Borra una tabla descargada y su registro de trazabilidad.

    Con `vacuum=True` además compacta el archivo: SQLite no devuelve el espacio al sistema
    al borrar una tabla, así que sin esto la base sigue ocupando los GB de un dataset grande.
    """
    try:
        _validate_table_name(table_name)
    except ValueError as e:
        return f"Error: {e}"

    try:
        with _db() as conn:
            existe = conn.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name = ?", (table_name,)
            ).fetchone()
            if not existe:
                return f"La tabla '{table_name}' no existe."

            filas = conn.execute(f'SELECT COUNT(*) FROM {_quote_ident(table_name)}').fetchone()[0]
            antes = os.path.getsize(DB_PATH) if os.path.exists(DB_PATH) else 0

            conn.execute(f'DROP TABLE {_quote_ident(table_name)}')
            try:
                conn.execute("DELETE FROM _meta_downloads WHERE table_name = ?", (table_name,))
            except sqlite3.OperationalError:
                pass  # nunca se registró una descarga para esta tabla
            conn.commit()

            if vacuum:
                # VACUUM no puede correr dentro de una transacción.
                conn.isolation_level = None
                conn.execute("VACUUM")
    except sqlite3.Error as e:
        return f"Error al borrar la tabla: {e}"

    despues = os.path.getsize(DB_PATH) if os.path.exists(DB_PATH) else 0
    detalle = f" La base pasó de {_human_size(antes)} a {_human_size(despues)}." if vacuum else (
        " El archivo no se encoge hasta que compactes con vacuum=True."
    )
    return f"Tabla '{table_name}' borrada ({filas} filas).{detalle}"


def mef_dataset_summary(table_name: str) -> str:
    try:
        _validate_table_name(table_name)
    except ValueError as e:
        return f"Error: {e}"

    try:
        with _db() as conn:
            cursor = conn.cursor()
            quoted = _quote_ident(table_name)
            cursor.execute(f'PRAGMA table_info({quoted})')
            cols = cursor.fetchall()
            if not cols:
                return f"La tabla '{table_name}' no existe."

            cursor.execute(f'SELECT COUNT(*) FROM {quoted}')
            total = cursor.fetchone()[0]

            # Todas las estadísticas en UNA sola pasada por la tabla. Una consulta por
            # columna recorría el archivo tantas veces como columnas: sobre un millón de
            # filas y 20 columnas eran 18,5 s frente a 9,7 s así.
            piezas: list[str] = []
            plan: list[tuple[str, str, bool]] = []
            for col in cols:
                name, ctype = col[1], col[2]
                qcol = _quote_ident(name)
                numerica = ctype in ("INTEGER", "REAL")
                piezas.append(f'SUM({qcol} IS NULL)')
                piezas.append(f'COUNT(DISTINCT {qcol})')
                if numerica:
                    piezas += [f'MIN({qcol})', f'MAX({qcol})', f'AVG({qcol})']
                plan.append((name, ctype, numerica))

            cursor.execute(f'SELECT {", ".join(piezas)} FROM {quoted}')
            valores = list(cursor.fetchone())

            summary = {"tabla": table_name, "total_filas": total, "columnas": {}}
            for name, ctype, numerica in plan:
                col_info: dict = {"tipo": ctype}
                col_info["nulos"] = valores.pop(0) or 0
                col_info["valores_distintos"] = valores.pop(0)
                if numerica:
                    col_info["min"] = valores.pop(0)
                    col_info["max"] = valores.pop(0)
                    avg = valores.pop(0)
                    col_info["promedio"] = round(avg, 2) if avg is not None else None
                summary["columnas"][name] = col_info

        return json.dumps(summary, indent=2, ensure_ascii=False)
    except Exception as e:
        return f"Error al generar el resumen: {e}"


def mef_generate_chart(query: str, chart_type: str, title: str) -> str:
    if chart_type not in CHART_TYPES:
        return f"chart_type debe ser uno de: {', '.join(CHART_TYPES)}."

    series = _chart_series(query)
    if isinstance(series, str):
        return series
    labels, values, nota = series

    chart_config = {
        "type": chart_type,
        "data": {"labels": labels, "datasets": [{"label": title, "data": values}]},
        "options": {"title": {"display": True, "text": title}},
    }

    url = f"https://quickchart.io/chart?c={urllib.parse.quote(json.dumps(chart_config))}"
    # Muchos servidores y navegadores cortan las URLs alrededor de los 8 KB; con demasiadas
    # filas el gráfico devolvería un 414 en vez de una imagen.
    if len(url) > 6000:
        return (
            f"La consulta genera {len(labels)} puntos: la URL de QuickChart quedaría demasiado larga. "
            "Agrega un LIMIT/GROUP BY a la consulta o usa mef_generate_chart_local."
        )
    return f"Gráfico generado con éxito{nota}. URL: {url}"


def mef_generate_chart_local(query: str, chart_type: str, title: str, filename: str = "") -> str:
    if chart_type not in CHART_TYPES:
        return f"chart_type debe ser uno de: {', '.join(CHART_TYPES)}."

    series = _chart_series(query)
    if isinstance(series, str):
        return series
    labels, values, nota = series

    fig, ax = plt.subplots(figsize=(8, 5))
    try:
        if chart_type == "bar":
            ax.bar(labels, values, color="#3b82f6")
            plt.xticks(rotation=45, ha="right")
        elif chart_type == "pie":
            ax.pie(values, labels=labels, autopct="%1.1f%%")
        else:
            ax.plot(labels, values, marker="o", color="#3b82f6")
            plt.xticks(rotation=45, ha="right")

        ax.set_title(title)
        fig.tight_layout()
        path = _output_path(filename or title, ".png", default="grafico")
        fig.savefig(path, dpi=150)
    finally:
        plt.close(fig)

    return f"Gráfico generado: {path}{nota}"


def mef_generate_report(
    query: str,
    title: str,
    formats: list[str] | None = None,
    filename_base: str = "",
    max_rows: int = MAX_REPORT_ROWS,
) -> str:
    if formats is None:
        formats = ["markdown"]

    if not _is_read_query(query):
        return "Error: mef_generate_report solo acepta consultas de lectura (SELECT / WITH)."

    valid_formats = ("markdown", "html", "pdf")
    formats = [f.strip().lower() for f in formats]
    unknown = [f for f in formats if f not in valid_formats]
    formats = [f for f in valid_formats if f in formats]  # orden estable, sin duplicados
    if not formats:
        extra = f" Formato(s) no reconocido(s): {', '.join(unknown)}." if unknown else ""
        return f"Error: 'formats' debe incluir al menos uno de: markdown, pdf, html.{extra}"

    try:
        # Se pide una fila de más para poder decir que había continuación sin traerse el
        # resto: un SELECT sin LIMIT sobre una tabla de millones de filas no cabe en RAM.
        columns, rows, truncated, source_info = _run_read_query(
            query, max_rows if max_rows > 0 else None
        )
    except sqlite3.Error as e:
        return f"Error al ejecutar la consulta: {e}"

    if not rows:
        return "La consulta no retornó filas; no se generó ningún informe."

    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    base = filename_base or title
    outputs = []

    try:
        if "markdown" in formats:
            path = _output_path(base, ".md")
            content = _build_markdown_report(title, columns, rows, generated_at, source_info, truncated)
            Path(path).write_text(content, encoding="utf-8")
            outputs.append(path)

        if "html" in formats:
            path = _output_path(base, ".html")
            content = _build_html_report(title, columns, rows, generated_at, source_info, truncated)
            Path(path).write_text(content, encoding="utf-8")
            outputs.append(path)

        if "pdf" in formats:
            path = _output_path(base, ".pdf")
            _build_pdf_report(title, columns, rows, generated_at, source_info, path, truncated)
            outputs.append(path)
    except OSError as e:
        return f"Error al escribir el informe: {e}"

    note = " (la consulta devuelve más filas: sube max_rows si las necesitas)" if truncated else ""
    return f"Informe generado ({len(rows)} filas{note}):\n" + "\n".join(outputs)


def mef_export_csv(query: str, filename: str) -> str:
    if not _is_read_query(query):
        return "Error: mef_export_csv solo acepta consultas de lectura (SELECT / WITH)."

    filepath = _output_path(filename, ".csv")
    escritas = 0
    try:
        # Se escribe por lotes según llegan las filas: exportar un millón de filas con
        # fetchall() pedía 1,2 GB de RAM, y aquí la memoria se mantiene plana.
        with _stream_read_query(query) as (columns, filas):
            with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
                writer = csv.writer(f)
                writer.writerow(columns)
                for fila in filas:
                    writer.writerow(fila)
                    escritas += 1
    except sqlite3.Error as e:
        return f"Error al ejecutar la consulta: {e}"
    except OSError as e:
        return f"Error al escribir el CSV: {e}"

    if not escritas:
        os.remove(filepath)
        return "La consulta no retornó filas."
    return f"Éxito. Datos exportados correctamente a {filepath} ({escritas} filas)."


def mef_export_excel(query: str, filename: str = "") -> str:
    if not _is_read_query(query):
        return "Error: mef_export_excel solo acepta consultas de lectura (SELECT / WITH)."

    path = _output_path(filename, ".xlsx")
    escritas = 0
    truncado = False

    try:
        with _stream_read_query(query) as (columns, filas):
            if not columns:
                return "La consulta no retornó columnas."

            wb = Workbook(write_only=True)
            ws = wb.create_sheet("Datos")

            # El ancho se calcula con la primera muestra que llega: medir el texto de un
            # millón de filas multiplicaría el tiempo sin cambiar el resultado visible.
            muestra = []
            for fila in filas:
                muestra.append(fila)
                if len(muestra) >= 500:
                    break
            for i, col in enumerate(columns, start=1):
                ancho = max([len(str(col))] + [len(str(f[i - 1])) for f in muestra if f[i - 1] is not None])
                ws.column_dimensions[get_column_letter(i)].width = min(ancho + 2, 50)

            ws.freeze_panes = "A2"
            cabecera = []
            for col in columns:
                celda = WriteOnlyCell(ws, value=col)
                celda.font = Font(bold=True)
                cabecera.append(celda)
            ws.append(cabecera)

            for fila in itertools.chain(muestra, filas):
                if escritas >= EXCEL_MAX_ROWS:
                    truncado = True
                    break
                ws.append(list(fila))
                escritas += 1
    except sqlite3.Error as e:
        return f"Error al ejecutar la consulta: {e}"

    if not escritas:
        return "La consulta no retornó filas."

    try:
        wb.save(path)
    except OSError as e:
        return f"Error al escribir el Excel: {e}"

    aviso = ""
    if truncado:
        aviso = (
            f" El formato .xlsx no admite más de {EXCEL_MAX_ROWS} filas de datos: el resto se "
            "omitió (usa mef_export_csv para el volcado completo)."
        )
    return f"Éxito. Datos exportados a {path} ({escritas} filas).{aviso}"
