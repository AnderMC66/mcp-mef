import json
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from mcp_mef import core

OUTPUT_FILE = os.path.join(core.OUTPUT_DIR, "cuadros_presupuesto_salud_2026.xlsx")

wb = Workbook()
wb.remove(wb.active)  # Quitar hoja por defecto

thin_border = Border(
    left=Side(style='thin', color='DDDDDD'),
    right=Side(style='thin', color='DDDDDD'),
    top=Side(style='thin', color='DDDDDD'),
    bottom=Side(style='thin', color='DDDDDD')
)

def agregar_hoja(wb, title, headers, rows, total_row=None):
    ws = wb.create_sheet(title=title)
    
    # Encabezado institucional
    ws.merge_cells("A1:H1")
    top = ws["A1"]
    top.value = f"MEF - CONSULTA AMIGABLE: {title.upper()}"
    top.font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    top.fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    top.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:H2")
    sub = ws["A2"]
    sub.value = "Datos oficiales del portal MEF Transparencia Económica / Consulta Amigable (Año 2026 - al 08 de octubre)"
    sub.font = Font(name="Segoe UI", size=9, italic=True, color="555555")
    sub.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 18

    # Cabeceras
    ws.append([])
    ws.append(headers)
    h_row = 4
    ws.row_dimensions[h_row].height = 22

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=h_row, column=col_idx)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="244061", end_color="244061", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    num_fmt = "#,##0.00"
    for r in rows:
        ws.append(r)
        curr = ws.max_row
        ws.row_dimensions[curr].height = 19
        for c in range(1, len(r) + 1):
            cell = ws.cell(row=curr, column=c)
            cell.font = Font(name="Segoe UI", size=9)
            cell.border = thin_border
            if c > 1 and c < len(r):
                cell.number_format = num_fmt
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c == len(r):
                cell.number_format = "0.0%"
                if isinstance(cell.value, (int, float)):
                    cell.value = cell.value / 100.0
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    if total_row:
        ws.append(total_row)
        t_row = ws.max_row
        ws.row_dimensions[t_row].height = 22
        for c in range(1, len(total_row) + 1):
            cell = ws.cell(row=t_row, column=c)
            cell.font = Font(name="Segoe UI", size=10, bold=True)
            cell.fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
            cell.border = thin_border
            if c > 1 and c < len(total_row):
                cell.number_format = num_fmt
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c == len(total_row):
                cell.number_format = "0.0%"
                if isinstance(cell.value, (int, float)):
                    cell.value = cell.value / 100.0
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for r_idx in range(4, ws.max_row + 1):
            v = ws.cell(row=r_idx, column=col_idx).value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width, 42)
    ws.freeze_panes = "A5"

# Datos para cada cuadro
headers = ["Descripción / Clasificador", "PIA (S/)", "PIM (S/)", "Certificación (S/)", "Comp. Anual (S/)", "Devengado (S/)", "Girado (S/)", "% Avance"]

# Cuadro 1: Salud en el Presupuesto Nacional
c1_rows = [
    ["TOTAL PRESUPUESTO NACIONAL", 257561619143.0, 288623963096.0, 262963972329.0, 243955018368.0, 196498129665.0, 193157545381.0, 68.1],
    ["20: FUNCIÓN SALUD (TODO EL ESTADO)", 32887683068.0, 35849839028.0, 32442283896.0, 30551756909.0, 25092232930.0, 24731852905.0, 70.0],
    ["011: MINISTERIO DE SALUD (MINSA CENTRAL)", 11396612777.0, 11308188481.0, 10032934508.0, 9752016645.0, 7829376100.0, 7770743613.0, 69.2]
]
agregar_hoja(wb, "1. Salud en Ppto Nacional", headers, c1_rows)

# Cuadro 2: Función Salud por Nivel de Gobierno
c2_rows = [
    ["E: GOBIERNO NACIONAL", 16349450831.0, 16415397311.0, 14964294547.0, 14556726289.0, 11889484264.0, 11810304440.0, 72.4],
    ["R: GOBIERNOS REGIONALES", 15750399090.0, 18204131791.0, 16486553037.0, 15199742917.0, 12560720772.0, 12295635702.0, 69.0],
    ["M: GOBIERNOS LOCALES", 787833147.0, 1230309926.0, 991436312.0, 795287702.0, 642027895.0, 625912763.0, 52.2]
]
c2_tot = ["TOTAL FUNCIÓN SALUD", 32887683068.0, 35849839028.0, 32442283896.0, 30551756909.0, 25092232930.0, 24731852905.0, 70.0]
agregar_hoja(wb, "2. Funcion Salud por Nivel", headers, c2_rows, c2_tot)

# Cuadro 3: Sector Salud por Pliego
c3_rows = [
    ["011: M. DE SALUD (MINSA)", 11396612777.0, 11308188481.0, 10032934508.0, 9752016645.0, 7829376100.0, 7770743613.0, 69.2],
    ["135: SEGURO INTEGRAL DE SALUD (SIS)", 2686113352.0, 2685895168.0, 2650008813.0, 2617695554.0, 2499897987.0, 2496039618.0, 93.1],
    ["136: INST. NAC. ENFERMEDADES NEOPLASICAS (INEN)", 351952428.0, 532854864.0, 511562921.0, 485448701.0, 374930167.0, 368849371.0, 70.4],
    ["131: INSTITUTO NACIONAL DE SALUD (INS)", 201932107.0, 202520374.0, 188992752.0, 170349959.0, 122582422.0, 120456799.0, 60.5],
    ["134: SUPERINTENDENCIA NACIONAL DE SALUD (SUSALUD)", 69844662.0, 69805009.0, 67273310.0, 66556218.0, 47597757.0, 47438650.0, 68.2]
]
c3_tot = ["TOTAL SECTOR SALUD", 14706455326.0, 14799263896.0, 13450772304.0, 13092067077.0, 10874384433.0, 10803528051.0, 73.5]
agregar_hoja(wb, "3. Sector Salud por Pliego", headers, c3_rows, c3_tot)

# Cuadro 4: MINSA por Fuente de Financiamiento
c4_rows = [
    ["00: RECURSOS ORDINARIOS", 11190571059.0, 10316207393.0, 9171283468.0, 8974024305.0, 7175947257.0, 7126861777.0, 69.6],
    ["13: DONACIONES Y TRANSFERENCIAS", 0.0, 729743483.0, 656245801.0, 600561491.0, 505056590.0, 496621549.0, 69.2],
    ["09: RECURSOS DIRECTAMENTE RECAUDADOS", 206041718.0, 222100903.0, 176341716.0, 148367328.0, 119880865.0, 118768899.0, 54.0],
    ["19: REC. POR OPERACIONES OFICIALES DE CREDITO", 0.0, 40000000.0, 29011847.0, 29011847.0, 28459201.0, 28459201.0, 71.1],
    ["18: CANON Y SOBRECANON, REGALIAS", 0.0, 136702.0, 51675.0, 51675.0, 32187.0, 32187.0, 23.5]
]
c4_tot = ["TOTAL MINSA", 11396612777.0, 11308188481.0, 10032934508.0, 9752016645.0, 7829376100.0, 7770743613.0, 69.2]
agregar_hoja(wb, "4. MINSA por Fuente", headers, c4_rows, c4_tot)

# Cuadro 5: MINSA por Genérica de Gasto
c5_rows = [
    ["5-21: PERSONAL Y OBLIGACIONES SOCIALES", 6334443545.0, 5454781834.0, 5125030805.0, 5059776514.0, 3850227350.0, 3826107114.0, 70.6],
    ["5-23: BIENES Y SERVICIOS", 2596098674.0, 3320456560.0, 2968800942.0, 2772515714.0, 2237311921.0, 2215250266.0, 67.4],
    ["6-26: ADQUISICION DE ACTIVOS NO FINANCIEROS", 1591509920.0, 1632410184.0, 1587875778.0, 1577446047.0, 1405876011.0, 1393521893.0, 86.1],
    ["5-24: DONACIONES Y TRANSFERENCIAS", 584576903.0, 691651282.0, 161251875.0, 152451997.0, 152385300.0, 152385300.0, 22.0],
    ["5-25: OTROS GASTOS", 160863863.0, 117095289.0, 101101092.0, 100970996.0, 96955100.0, 96948295.0, 82.8],
    ["5-22: PENSIONES Y OTRAS PRESTACIONES SOCIALES", 129119872.0, 90793332.0, 87874016.0, 87855377.0, 85620419.0, 85530746.0, 94.3],
    ["6-24: DONACIONES Y TRANSFERENCIAS CAPITAL", 0.0, 1000000.0, 1000000.0, 1000000.0, 1000000.0, 1000000.0, 100.0]
]
c5_tot = ["TOTAL MINSA", 11396612777.0, 11308188481.0, 10032934508.0, 9752016645.0, 7829376100.0, 7770743613.0, 69.2]
agregar_hoja(wb, "5. MINSA por Generica", headers, c5_rows, c5_tot)

# Cuadro 6: MINSA por Programa Presupuestal
c6_rows = [
    ["9002: ASIGNACIONES PPTALES QUE NO RESULTAN EN PRODUCTOS", 5378081454.0, 5693544718.0, 5249213775.0, 5068619266.0, 4149995334.0, 4110218684.0, 72.9],
    ["9001: ACCIONES CENTRALES", 3758588615.0, 2585741212.0, 2359615447.0, 2321010742.0, 1748014976.0, 1741816725.0, 67.6],
    ["0104: REDUCCION DE MORTALIDAD POR EMERGENCIAS Y URGENCIAS", 280360696.0, 516924701.0, 501188508.0, 497387380.0, 415252378.0, 412119212.0, 80.3],
    ["0002: SALUD MATERNO NEONATAL", 399906097.0, 557343508.0, 484316352.0, 472365484.0, 397220518.0, 393948108.0, 71.3],
    ["0024: PREVENCION Y CONTROL DEL CANCER", 495321079.0, 558216220.0, 464675957.0, 440440646.0, 365448351.0, 364670049.0, 65.5],
    ["1001: PRODUCTOS ESPECIFICOS PARA DESARROLLO INFANTIL TEMPRANO", 440715399.0, 525453266.0, 269788783.0, 265223086.0, 212268541.0, 210992250.0, 40.4],
    ["0018: ENFERMEDADES NO TRANSMISIBLES", 112798380.0, 198240593.0, 183359366.0, 181001636.0, 147124199.0, 145725646.0, 74.2],
    ["0016: TBC-VIH/SIDA", 250841875.0, 285188367.0, 204858479.0, 197279606.0, 143702594.0, 142449374.0, 50.4],
    ["0131: CONTROL Y PREVENCION EN SALUD MENTAL", 113699227.0, 174053234.0, 146393778.0, 142304978.0, 113005990.0, 112285306.0, 64.9],
    ["0129: PREVENCION Y MANEJO EN DISCAPACIDAD", 63444553.0, 86462994.0, 83202908.0, 81041874.0, 66992039.0, 66828552.0, 77.5],
    ["0017: ENFERMEDADES METAXENICAS Y ZOONOSIS", 73783053.0, 86750620.0, 50398605.0, 49799464.0, 41670673.0, 41169722.0, 48.0],
    ["0068: REDUCCION DE VULNERABILIDAD ANTE DESASTRES", 15020454.0, 23318919.0, 19613831.0, 19241765.0, 15610835.0, 15459367.0, 66.9],
    ["1002: REDUCCION DE VIOLENCIA CONTRA LA MUJER", 14051895.0, 16950129.0, 16308718.0, 16300718.0, 13069670.0, 13060619.0, 77.1]
]
c6_tot = ["TOTAL MINSA", 11396612777.0, 11308188481.0, 10032934508.0, 9752016645.0, 7829376100.0, 7770743613.0, 69.2]
agregar_hoja(wb, "6. MINSA por Programa", headers, c6_rows, c6_tot)

wb.save(OUTPUT_FILE)
print("Archivo generado exitosamente en:", OUTPUT_FILE)
