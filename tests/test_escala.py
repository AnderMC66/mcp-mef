"""Comportamiento sobre tablas grandes: topes, streaming y avisos.

Los datasets reales del MEF llegan a 7,7 millones de filas; estas pruebas fijan las
salvaguardas que evitan que una consulta inocente devuelva cientos de MB.
"""
import json
import os

import pytest
from openpyxl import load_workbook

from mcp_mef import core


@pytest.fixture()
def muchas_filas(db):
    """5 000 filas: suficiente para cruzar los topes sin ralentizar la suite."""
    with db._db() as conn:
        conn.execute("CREATE TABLE grande (id INTEGER, dep TEXT, monto REAL)")
        conn.executemany(
            "INSERT INTO grande VALUES (?, ?, ?)",
            [(i, f"DEP_{i % 7}", float(i)) for i in range(5000)],
        )
        conn.commit()
    return db


def test_sql_corta_en_max_rows(muchas_filas):
    payload = json.loads(core.sql_mef_db("SELECT * FROM grande", max_rows=10))
    assert len(payload["filas"]) == 10
    assert payload["truncado"] is True
    assert "LIMIT" in payload["nota"]


def test_sql_sin_truncar_no_avisa(muchas_filas):
    payload = json.loads(core.sql_mef_db("SELECT * FROM grande LIMIT 5", max_rows=10))
    assert len(payload["filas"]) == 5
    assert "truncado" not in payload


def test_sql_max_rows_cero_trae_todo(muchas_filas):
    payload = json.loads(core.sql_mef_db("SELECT id FROM grande", max_rows=0))
    assert len(payload["filas"]) == 5000


def test_sql_tope_por_defecto(muchas_filas):
    payload = json.loads(core.sql_mef_db("SELECT * FROM grande"))
    assert len(payload["filas"]) == core.MAX_SQL_ROWS


def test_grafico_rechaza_demasiados_puntos(muchas_filas):
    salida = core.mef_generate_chart_local("SELECT id, monto FROM grande", "bar", "x")
    assert str(core.MAX_CHART_POINTS) in salida
    assert "GROUP BY" in salida


def test_grafico_agrupado_si_cabe(muchas_filas):
    salida = core.mef_generate_chart_local(
        "SELECT dep, SUM(monto) FROM grande GROUP BY dep", "bar", "Por dep"
    )
    assert salida.startswith("Gráfico generado:")


def test_csv_exporta_todo_sin_tope(muchas_filas):
    """El CSV es el volcado completo: no debe heredar el tope de sql_mef_db."""
    core.mef_export_csv("SELECT id FROM grande", "todo")
    ruta = os.path.join(core.OUTPUT_DIR, "todo.csv")
    with open(ruta, encoding="utf-8-sig") as f:
        lineas = sum(1 for _ in f)
    assert lineas == 5001  # cabecera + filas


def test_csv_vacio_no_deja_archivo(muchas_filas):
    salida = core.mef_export_csv("SELECT id FROM grande WHERE id < 0", "vacio")
    assert "no retornó filas" in salida
    assert not os.path.exists(os.path.join(core.OUTPUT_DIR, "vacio.csv"))


def test_excel_conserva_formato_y_contenido(muchas_filas):
    core.mef_export_excel("SELECT id, dep FROM grande LIMIT 3", "libro")
    wb = load_workbook(os.path.join(core.OUTPUT_DIR, "libro.xlsx"))
    ws = wb["Datos"]
    assert [c.value for c in ws[1]] == ["id", "dep"]
    assert ws[1][0].font.bold
    assert ws.freeze_panes == "A2"
    assert ws.max_row == 4
    assert ws.column_dimensions["B"].width > 0


def test_excel_exporta_muchas_filas(muchas_filas):
    salida = core.mef_export_excel("SELECT id FROM grande", "grande")
    assert "5000 filas" in salida
    wb = load_workbook(os.path.join(core.OUTPUT_DIR, "grande.xlsx"), read_only=True)
    # En modo write_only openpyxl no graba la dimensión de la hoja, así que se cuentan
    # las filas recorriéndolas (Excel las abre igual).
    assert sum(1 for _ in wb["Datos"].iter_rows()) == 5001


def test_resumen_de_columna_numerica_y_texto(muchas_filas):
    resumen = json.loads(core.mef_dataset_summary("grande"))
    assert resumen["total_filas"] == 5000
    assert resumen["columnas"]["monto"]["max"] == 4999.0
    assert resumen["columnas"]["monto"]["promedio"] == 2499.5
    assert resumen["columnas"]["dep"]["valores_distintos"] == 7
    assert "min" not in resumen["columnas"]["dep"]
